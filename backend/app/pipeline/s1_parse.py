"""s1 Parser (Reader): 从 MinIO / 本地读取 docx 并解析为 raw_text / raw_tables / raw_images。"""

from __future__ import annotations

import re
from pathlib import Path
from urllib.parse import urlparse

import structlog
from docx import Document
from docx.oxml import CT_P, CT_Tbl
from docx.oxml.ns import qn

from app.pipeline.base import stage
from app.pipeline.context import PipelineContext
from app.services.doc_converter import open_as_docx
from app.services.doc_format import INVALID_DOCX_MSG, read_sniff_word_format
from app.services.legacy_doc_text import extract_text_from_legacy_doc
from app.services.minio_client import download_to_tempfile

logger = structlog.get_logger(__name__)

HEADING_RE = re.compile(r"^Heading(\d+)$", re.IGNORECASE)


def _style_level_from_xml(p_elem) -> int | None:
    """直接从 CT_P XML 读 pStyle val, 避免 python-docx style.name 失效问题。"""
    p_pr = p_elem.find(qn("w:pPr"))
    if p_pr is None:
        return None
    p_style = p_pr.find(qn("w:pStyle"))
    if p_style is None:
        return None
    val = p_style.get(qn("w:val")) or ""
    m = HEADING_RE.match(val.strip())
    return int(m.group(1)) if m else None


def _para_text(p_elem) -> str:
    """从 CT_P 提取所有 w:t 文本。"""
    texts = p_elem.xpath(".//w:t")
    return "".join(t.text or "" for t in texts).strip()


def _load_local_docx(path: str) -> Document:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Local docx not found: {p}")
    if read_sniff_word_format(p) != "docx":
        raise ValueError(INVALID_DOCX_MSG)
    return Document(str(p))


def _read_doc_source(path: Path) -> tuple[Document | None, str | None]:
    """优先返回可供 python-docx 解析的 Document; 失败时回退 legacy doc 文本。"""
    fmt = read_sniff_word_format(path)
    if fmt == "docx":
        with open_as_docx(path) as docx_path:
            return _load_local_docx(str(docx_path)), None
    if fmt == "doc":
        try:
            with open_as_docx(path) as docx_path:
                return _load_local_docx(str(docx_path)), None
        except RuntimeError as exc:
            fallback_text = extract_text_from_legacy_doc(path)
            if not fallback_text:
                raise RuntimeError(f"legacy .doc 文本兜底解析失败: {path}") from exc
            logger.warning(
                "s1_parse.legacy_doc_fallback",
                source=str(path),
                reason=str(exc),
                text_len=len(fallback_text),
            )
            return None, fallback_text
    raise ValueError(INVALID_DOCX_MSG)


def _tbl_to_dict(table, section_path: list[str], raw_index: int) -> dict:
    hdr = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
    rows = []
    for row in table.rows[1:]:
        rows.append(dict(zip(hdr, [c.text.strip() for c in row.cells], strict=False)))
    return {
        "section_path": section_path.copy(),
        "headers": hdr,
        "rows": rows,
        "raw_index": raw_index,
    }


_REPORT_ID_HEADER_KEYS = ("项目编号", "报告编号", "Report ID", "report_id")
_REPORT_ID_TEXT_RE = re.compile(
    r"(?:FS|8D|RPT)[-_]?[A-Z0-9][A-Z0-9\-_/]{3,}",
    re.IGNORECASE,
)


def _extract_report_id_from_tables(raw_tables: list[dict]) -> str | None:
    """扫 raw_tables 第一个含『项目编号/报告编号』表头的表, 取第一行该列值。

    注意: `_tbl_to_dict` 返回 rows=[data_row_dict, ...] (不含 header 行),
    而 headers=[col_name, ...]。当 rows 仅含 1 行时, 数据行即 rows[0]。
    """
    for tbl in raw_tables:
        headers = tbl.get("headers", [])
        rows = tbl.get("rows", [])
        # 找含项目编号/报告编号的列索引
        col_idx = None
        for idx, h in enumerate(headers):
            if any(key in h for key in _REPORT_ID_HEADER_KEYS):
                col_idx = idx
                break
        if col_idx is None:
            continue
        # 取第一行数据的该列值
        if rows:
            first_row = rows[0]
            if isinstance(first_row, dict) and col_idx < len(headers):
                header_key = headers[col_idx]
                value = str(first_row.get(header_key, "")).strip()
                if value:
                    return value
    return None


def _extract_report_id_from_text(raw_text: str) -> str | None:
    """从正文前 2000 字匹配 FS-/8D- 样式报告编号。"""
    if not raw_text:
        return None
    m = _REPORT_ID_TEXT_RE.search(raw_text[:2000])
    if m:
        return m.group(0).upper().replace("_", "-")
    return None


@stage("s1_parse")
async def run(ctx: PipelineContext) -> PipelineContext:
    """从 ctx.minio_key 读取 docx, 解析为 raw_text / raw_tables / raw_images。"""
    parsed = urlparse(ctx.minio_key)
    scheme = parsed.scheme

    if scheme == "local":
        local_path = Path(ctx.minio_key[len("local://") :])
        doc, fallback_text = _read_doc_source(local_path)
    elif scheme == "minio":
        tmp_path = await download_to_tempfile(ctx.minio_key)
        try:
            doc, fallback_text = _read_doc_source(tmp_path)
        finally:
            tmp_path.unlink(missing_ok=True)
    else:
        raise ValueError(f"Unsupported minio_key scheme: {scheme}")

    if fallback_text is not None:
        ctx.raw_text = fallback_text
        ctx.raw_tables = []
        ctx.raw_images = []

        if not ctx.report_id_hint:
            extracted_id = _extract_report_id_from_text(ctx.raw_text)
            if extracted_id:
                ctx.report_id_hint = extracted_id
                logger.info("s1_parse.report_id_extracted", report_id=extracted_id)

        logger.info(
            "s1_parse.done",
            text_len=len(ctx.raw_text),
            tables=0,
            images=0,
            report_id_hint=ctx.report_id_hint,
            reader_backend="legacy_doc_text_fallback",
        )
        return ctx

    assert doc is not None

    section_path: list[str] = []
    raw_tables: list[dict] = []
    raw_text_parts: list[str] = []
    last_sec_marker = ""

    body = doc.element.body
    t_idx = 0  # doc.tables index

    for child in body:
        if isinstance(child, CT_P):
            level = _style_level_from_xml(child)
            if level is not None:
                heading_text = _para_text(child)
                if heading_text:
                    section_path = [*section_path[: level - 1], heading_text]
                    marker = "\n[SEC:" + "/".join(section_path) + "]\n\n"
                    if marker != last_sec_marker:
                        raw_text_parts.append(marker)
                        last_sec_marker = marker
            else:
                text = _para_text(child)
                if text:
                    raw_text_parts.append(text + "\n\n")
        elif isinstance(child, CT_Tbl):
            if t_idx < len(doc.tables):
                tbl_dict = _tbl_to_dict(doc.tables[t_idx], section_path, t_idx)
                raw_tables.append(tbl_dict)
            t_idx += 1

    ctx.raw_text = "".join(raw_text_parts).strip()
    ctx.raw_tables = raw_tables
    ctx.raw_images = []

    # --- 提取 report_id (如果调用方没传 hint) ---
    if not ctx.report_id_hint:
        extracted_id = _extract_report_id_from_tables(ctx.raw_tables)
        if not extracted_id:
            extracted_id = _extract_report_id_from_text(ctx.raw_text)
        if extracted_id:
            ctx.report_id_hint = extracted_id
            logger.info("s1_parse.report_id_extracted", report_id=extracted_id)

    logger.info(
        "s1_parse.done",
        text_len=len(ctx.raw_text),
        tables=len(ctx.raw_tables),
        images=len(ctx.raw_images),
        report_id_hint=ctx.report_id_hint,
    )
    return ctx
