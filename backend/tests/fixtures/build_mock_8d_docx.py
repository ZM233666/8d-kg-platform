"""生成最小 8D mock 报告 docx，用于端到端测试。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def _add_paragraph(doc: Document, text: str):
    return doc.add_paragraph(text)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    # Header row
    hdr_row = tbl.rows[0]
    for j, h in enumerate(headers):
        cell = hdr_row.cells[j]
        cell.text = h
        run = cell.paragraphs[0].runs[0]
        run.bold = True
    # Data rows
    for i, row_data in enumerate(rows):
        row = tbl.rows[i + 1]
        for j, val in enumerate(row_data):
            row.cells[j].text = val
    return tbl


def build() -> Path:
    doc = Document()
    # D1 项目信息
    _add_heading(doc, "D1 项目信息", level=1)
    _add_table(
        doc,
        headers=["项目编号", "客户", "车型", "报告编号"],
        rows=[["FS-2024-001", "佛山南海地铁", "AW3", "RPT-001"]],
    )

    # D2 故障描述
    _add_heading(doc, "D2 故障描述", level=1)
    _add_table(
        doc,
        headers=["序号", "故障描述", "位置", "里程"],
        rows=[["1", "EP2002阀二级调节器弹簧端部断裂", "靠近端部第二圈", "125000 km"]],
    )

    # D4 根本原因分析
    _add_heading(doc, "D4 根本原因分析", level=1)
    _add_table(
        doc,
        headers=["项目", "目标值", "上限", "下限", "实测值", "单位", "判定"],
        rows=[["硬度 HV10", "480", "520", "440", "475", "HV", "合格"]],
    )

    # D5 纠正措施（3段：containment / corrective / preventive）
    _add_heading(doc, "D5 纠正措施", level=1)
    _add_paragraph(
        doc,
        "【临时措施】对在用车辆全部 EP2002 阀进行 100% 检查，发现异常立即更换。责任人：张三。",
    )
    doc.add_paragraph("")  # 空行分隔
    _add_paragraph(
        doc,
        "【永久措施】更换塑料活塞销供应商，启用 B 厂金属销方案，从根本上消除气孔缺陷。责任人：李四。",
    )
    doc.add_paragraph("")  # 空行分隔
    _add_paragraph(
        doc,
        "【预防措施】建立活塞销来料 100% X-ray 检测工序，防止不合格品流入装配线。责任人：王五。",
    )

    # D8 结案
    _add_heading(doc, "D8 结案", level=1)
    _add_paragraph(
        doc,
        "经团队评审，根因已确认，临时措施已实施，永久措施正在推进中，风险可控，建议结案关闭。",
    )

    out_path = Path(__file__).parent / "mock_8d_report.docx"
    doc.save(str(out_path))
    return out_path


if __name__ == "__main__":
    p = build()
    print(p)
